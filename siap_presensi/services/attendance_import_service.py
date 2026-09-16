"""
Modul Layanan Import Absensi Excel (Attendance Import Service) SIAP.
Mengkoordinasikan seluruh alur kerja Tahap 3:
1. Membaca berkas Excel .xlsx dan daftar worksheet.
2. Memetakan header otomatis dan memvalidasi baris data.
3. Eksekusi penyimpanan batch ke tabel attendance_raw dalam transaksi database yang aman (ACID).
4. Penanganan duplikasi (Lewati / Simpan Baru).
5. Penanganan karyawan tidak ditemukan (Lewati / Buat Master / Simpan Tanpa Relasi).
6. Pencatatan batch ke import_logs dan audit_logs.
7. Pembuatan dan ekspor Error Log (.xlsx / .csv).
8. Pengambilan riwayat import dan detail batch.
"""
import os
import json
from datetime import datetime, date
from pathlib import Path
from typing import Dict, Any, List, Optional, Tuple
from sqlalchemy.orm import Session
from sqlalchemy import desc

from database.models import AttendanceRaw, Employee, EmployeeStatus, ImportLog, AuditLog, User
from database.connection import get_db_session
from utils.excel_reader import get_excel_sheet_names, read_excel_rows
from utils.date_parser import format_date_display
from utils.time_parser import format_time_display
from utils.logger import get_logger
from services.import_validation_service import ImportValidationService, map_excel_headers

logger = get_logger("AttendanceImportService")


def generate_batch_id(db: Session) -> str:
    """
    Menghasilkan Batch ID berformat BATCH-YYYYMMDD-XXX (contoh: BATCH-20260916-001).
    """
    today_str = datetime.utcnow().strftime("%Y%m%d")
    prefix = f"BATCH-{today_str}-"

    # Cari batch terakhir pada hari yang sama
    last_batch = (
        db.query(ImportLog)
        .filter(ImportLog.import_batch_id.like(f"{prefix}%"))
        .order_by(desc(ImportLog.import_batch_id))
        .first()
    )

    if last_batch and last_batch.import_batch_id:
        try:
            last_seq = int(last_batch.import_batch_id.split("-")[-1])
            new_seq = last_seq + 1
        except Exception:
            new_seq = 1
    else:
        new_seq = 1

    return f"{prefix}{new_seq:03d}"


class AttendanceImportService:
    """Layanan orkestrator proses import absensi Excel."""

    @staticmethod
    def get_worksheets(filepath: str) -> List[str]:
        """Mengambil daftar worksheet yang dapat dipilih dari file Excel."""
        return get_excel_sheet_names(filepath)

    @staticmethod
    def preview_excel_file(filepath: str, sheet_name: Optional[str] = None) -> Dict[str, Any]:
        """
        Membaca dan memvalidasi isi berkas Excel tanpa menyimpan ke database.

        Returns:
            Dict berisi ringkasan metrik, header mapping, dan daftar baris terverifikasi.
        """
        path = Path(filepath)
        if not path.exists():
            raise FileNotFoundError(f"File '{filepath}' tidak ditemukan.")

        # 1. Baca data baris mentah
        headers, data_rows, total_count = read_excel_rows(filepath, sheet_name=sheet_name)
        if not headers or total_count == 0:
            raise ValueError("Worksheet Excel kosong atau tidak memiliki data baris.")

        # 2. Petakan header
        mapped_headers, missing_required, warnings = map_excel_headers(headers)
        if missing_required:
            missing_str = ", ".join(missing_required)
            raise ValueError(
                f"Format Excel tidak sesuai. Kolom {missing_str} wajib tersedia di lembar kerja."
            )

        # 3. Validasi baris demi baris menggunakan sesi database
        with get_db_session() as db:
            validator = ImportValidationService(db)
            validation_result = validator.validate_rows(
                raw_rows=data_rows,
                header_mapping=mapped_headers,
                source_filename=path.name,
            )
            validation_result["sheet_name"] = sheet_name or "ActiveSheet"
            validation_result["warnings"] = warnings
            return validation_result

    @staticmethod
    def execute_import(
        preview_data: Dict[str, Any],
        duplicate_strategy: str = "SKIP",     # "SKIP" atau "INSERT"
        unmatched_strategy: str = "UNLINKED", # "UNLINKED", "SKIP", atau "CREATE"
        user_id: Optional[int] = None,
        username: str = "SYSTEM",
    ) -> Dict[str, Any]:
        """
        Mengeksekusi import data ke database dalam transaksi ACID terlindungi.

        Args:
            preview_data: Hasil dari preview_excel_file()
            duplicate_strategy: Opsi penanganan duplikat ("SKIP" / "INSERT")
            unmatched_strategy: Opsi karyawan tidak ditemukan ("UNLINKED" / "SKIP" / "CREATE")
            user_id: ID pengguna yang mengeksekusi
            username: Username untuk audit trail

        Returns:
            Ringkasan hasil import (batch_id, success, failed, skipped, duplicates).
        """
        rows = preview_data.get("rows", [])
        source_filename = preview_data.get("source_file", "unknown.xlsx")
        total_rows = len(rows)

        if total_rows == 0:
            raise ValueError("Tidak ada baris data untuk di-import.")

        with get_db_session() as db:
            # 1. Buat Batch ID baru
            batch_id = generate_batch_id(db)

            # Inisialisasi catatan log status PROCESSING
            import_log = ImportLog(
                import_batch_id=batch_id,
                user_id=user_id,
                file_name=source_filename,
                import_type="EXCEL_ABSENSI",
                total_rows=total_rows,
                success_rows=0,
                failed_rows=0,
                duplicate_rows=0,
                status="PROCESSING",
                error_details="",
                created_at=datetime.utcnow(),
            )
            db.add(import_log)
            db.flush()

            success_count = 0
            failed_count = 0
            duplicate_count = 0
            skipped_count = 0
            error_records: List[Dict[str, Any]] = []

            try:
                for row in rows:
                    # A. Jika baris sejak awal tidak valid (error tanggal/jam/nama)
                    if not row.get("is_valid", False):
                        failed_count += 1
                        error_records.append({
                            "excel_line": row.get("excel_line", "-"),
                            "nama": row.get("nama", ""),
                            "tanggal": row.get("tanggal_raw", ""),
                            "scan_masuk": row.get("scan_masuk_raw", ""),
                            "scan_pulang": row.get("scan_pulang_raw", ""),
                            "status": "INVALID",
                            "error": "; ".join(row.get("errors", [])),
                        })
                        continue

                    # B. Penanganan Duplikasi
                    is_dup = row.get("is_duplicate_in_file") or row.get("is_duplicate_in_db")
                    if is_dup:
                        duplicate_count += 1
                        if duplicate_strategy == "SKIP":
                            skipped_count += 1
                            error_records.append({
                                "excel_line": row.get("excel_line", "-"),
                                "nama": row.get("nama", ""),
                                "tanggal": row.get("tanggal_display", ""),
                                "scan_masuk": row.get("scan_masuk_display", ""),
                                "scan_pulang": row.get("scan_pulang_display", ""),
                                "status": "DUPLIKAT_DILEWATI",
                                "error": row.get("duplicate_message", "Transaksi duplikat dilewati."),
                            })
                            continue

                    # C. Penanganan Karyawan Tidak Ditemukan
                    emp_id = row.get("employee_id")
                    if emp_id is None:
                        if unmatched_strategy == "SKIP":
                            skipped_count += 1
                            error_records.append({
                                "excel_line": row.get("excel_line", "-"),
                                "nama": row.get("nama", ""),
                                "tanggal": row.get("tanggal_display", ""),
                                "scan_masuk": row.get("scan_masuk_display", ""),
                                "scan_pulang": row.get("scan_pulang_display", ""),
                                "status": "KARYAWAN_TIDAK_DITEMUKAN",
                                "error": "Karyawan tidak ada di Master Data (dilewati).",
                            })
                            continue
                        elif unmatched_strategy == "CREATE":
                            # Tambahkan karyawan baru ke Master Data
                            new_emp = Employee(
                                emp_num=row.get("emp_num") or None,
                                no_id=row.get("no_id") or None,
                                nik=row.get("nik") or None,
                                nama=row.get("nama"),
                                status=EmployeeStatus.AKTIF,
                                keterangan="Dibuat otomatis dari Import Absensi Excel.",
                                created_at=datetime.utcnow(),
                                updated_at=datetime.utcnow(),
                            )
                            db.add(new_emp)
                            db.flush()
                            emp_id = new_emp.id

                    # D. Persistensi ke tabel attendance_raw
                    # PENTING: Data mentah disimpan apa adanya tanpa kalkulasi alfa / hari kerja / potongan!
                    raw_record = AttendanceRaw(
                        employee_id=emp_id,
                        emp_num=row.get("emp_num") or None,
                        no_id=row.get("no_id") or None,
                        nik=row.get("nik") or None,
                        nama=row.get("nama", ""),
                        tanggal=row.get("tanggal_obj"),
                        scan_masuk=row.get("scan_masuk"),
                        scan_pulang=row.get("scan_pulang"),
                        source_file=source_filename,
                        import_batch_id=batch_id,
                        created_at=datetime.utcnow(),
                    )
                    db.add(raw_record)
                    success_count += 1

                # Tentukan status akhir import batch
                if failed_count == 0 and skipped_count == 0:
                    final_status = "SUCCESS"
                elif success_count > 0:
                    final_status = "PARTIAL"
                else:
                    final_status = "FAILED"

                # Update Import Log
                import_log.success_rows = success_count
                import_log.failed_rows = failed_count + skipped_count
                import_log.duplicate_rows = duplicate_count
                import_log.status = final_status
                import_log.error_details = json.dumps(error_records, ensure_ascii=False)

                # Catat Audit Trail
                audit = AuditLog(
                    user_id=user_id,
                    action="IMPORT_ATTENDANCE",
                    module="ATTENDANCE_RAW",
                    description=(
                        f"Import absensi file '{source_filename}' (Batch: {batch_id}). "
                        f"Total: {total_rows}, Berhasil: {success_count}, Gagal/Dilewati: {failed_count + skipped_count}, "
                        f"Duplikat: {duplicate_count}. Status: {final_status}."
                    ),
                    created_at=datetime.utcnow(),
                )
                db.add(audit)
                db.commit()

                logger.info(f"Import Batch {batch_id} selesai: {final_status} ({success_count}/{total_rows} baris).")

                return {
                    "batch_id": batch_id,
                    "status": final_status,
                    "file_name": source_filename,
                    "total_rows": total_rows,
                    "success_rows": success_count,
                    "failed_rows": failed_count,
                    "skipped_rows": skipped_count,
                    "duplicate_rows": duplicate_count,
                    "error_records": error_records,
                }

            except Exception as e:
                db.rollback()
                logger.error(f"Kegagalan kritis saat import batch: {e}", exc_info=True)
                # Tandai batch gagal jika memungkinkan
                try:
                    with get_db_session() as err_db:
                        err_log = err_db.query(ImportLog).filter(ImportLog.import_batch_id == batch_id).first()
                        if err_log:
                            err_log.status = "FAILED"
                            err_log.error_details = json.dumps([{"error": f"Database Error: {str(e)}"}])
                            err_db.commit()
                except Exception:
                    pass
                raise RuntimeError(f"Gagal memproses import data ke database: {str(e)}")

    @staticmethod
    def export_error_log(error_records: List[Dict[str, Any]], filepath: str) -> str:
        """
        Mengekspor daftar error / peringatan baris import ke file Excel atau CSV.

        Kolom:
        No Baris Excel | Nama | Tanggal | Scan Masuk | Scan Pulang | Status | Keterangan Error
        """
        headers = [
            "No Baris Excel", "Nama", "Tanggal", "Scan Masuk", "Scan Pulang", "Status", "Keterangan Error"
        ]

        try:
            import openpyxl
            from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Error Log Import"

            header_font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
            header_fill = PatternFill(start_color="991B1B", end_color="991B1B", fill_type="solid") # Merah Tua
            thin_border = Border(
                left=Side(style="thin", color="CBD5E1"),
                right=Side(style="thin", color="CBD5E1"),
                top=Side(style="thin", color="CBD5E1"),
                bottom=Side(style="thin", color="CBD5E1")
            )

            ws.append(headers)
            for c_idx in range(1, len(headers) + 1):
                cell = ws.cell(row=1, column=c_idx)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = Alignment(horizontal="center", vertical="center")

            for rec in error_records:
                row_vals = [
                    rec.get("excel_line", "-"),
                    rec.get("nama", "-"),
                    rec.get("tanggal", "-"),
                    rec.get("scan_masuk", "-"),
                    rec.get("scan_pulang", "-"),
                    rec.get("status", "-"),
                    rec.get("error", "-"),
                ]
                ws.append(row_vals)
                row_num = ws.max_row
                for c_idx in range(1, len(row_vals) + 1):
                    ws.cell(row=row_num, column=c_idx).border = thin_border

            # Adjust width
            for col in ws.columns:
                max_len = max(len(str(c.value or "")) for c in col)
                col_letter = openpyxl.utils.get_column_letter(col[0].column)
                ws.column_dimensions[col_letter].width = max(max_len + 3, 14)

            wb.save(filepath)
            return filepath

        except ImportError:
            import csv
            csv_path = filepath if filepath.endswith(".csv") else filepath.rsplit(".", 1)[0] + ".csv"
            with open(csv_path, "w", newline="", encoding="utf-8-sig") as f:
                writer = csv.writer(f, delimiter=";")
                writer.writerow(headers)
                for rec in error_records:
                    writer.writerow([
                        rec.get("excel_line", "-"),
                        rec.get("nama", "-"),
                        rec.get("tanggal", "-"),
                        rec.get("scan_masuk", "-"),
                        rec.get("scan_pulang", "-"),
                        rec.get("status", "-"),
                        rec.get("error", "-"),
                    ])
            return csv_path

    @staticmethod
    def get_import_history(
        db: Session,
        keyword: Optional[str] = None,
        status: Optional[str] = None,
        page: int = 1,
        page_size: int = 20,
    ) -> Tuple[List[Dict[str, Any]], int, int]:
        """
        Mengambil riwayat batch import untuk Halaman Riwayat Import.

        Returns:
            (history_items, total_records, total_pages)
        """
        query = db.query(ImportLog, User.full_name).outerjoin(User, ImportLog.user_id == User.id)

        if keyword and keyword.strip():
            kw = f"%{keyword.strip()}%"
            query = query.filter(
                ImportLog.file_name.ilike(kw) | ImportLog.import_batch_id.ilike(kw)
            )

        if status and status != "SEMUA":
            query = query.filter(ImportLog.status == status)

        total_records = query.count()
        total_pages = max(1, (total_records + page_size - 1) // page_size)
        offset = (page - 1) * page_size

        results = query.order_by(desc(ImportLog.created_at)).offset(offset).limit(page_size).all()

        history_items = []
        for idx, (log_item, user_name) in enumerate(results, start=offset + 1):
            d = log_item.to_dict()
            d["no"] = idx
            d["user_name"] = user_name or "System / Admin"
            # Cek perkiraan periode transaksi berdasarkan attendance_raw terkait
            period_sample = (
                db.query(AttendanceRaw.tanggal)
                .filter(AttendanceRaw.import_batch_id == log_item.import_batch_id)
                .first()
            )
            if period_sample and period_sample[0]:
                d["periode"] = period_sample[0].strftime("%B %Y")
            else:
                d["periode"] = log_item.created_at.strftime("%B %Y")

            history_items.append(d)

        return history_items, total_records, total_pages

    @staticmethod
    def get_batch_details(db: Session, batch_id: str) -> Optional[Dict[str, Any]]:
        """Mengambil detail lengkap satu batch import termasuk log error-nya."""
        log_item = db.query(ImportLog).filter(ImportLog.import_batch_id == batch_id).first()
        if not log_item:
            return None

        data = log_item.to_dict()
        try:
            data["errors_list"] = json.loads(log_item.error_details or "[]")
        except Exception:
            data["errors_list"] = []

        # Hitung catatan mentah yang tersimpan
        data["raw_count"] = (
            db.query(AttendanceRaw)
            .filter(AttendanceRaw.import_batch_id == batch_id)
            .count()
        )
        return data
