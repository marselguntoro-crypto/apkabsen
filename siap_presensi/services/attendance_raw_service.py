"""
Modul Layanan Data Absensi Mentah (Attendance Raw Service) SIAP.
Mengelola kueri, penyaringan multi-parameter, paginasi, statistik,
dan ekspor data transaksi absensi mentah dari tabel attendance_raw.

PENTING: Modul ini HANYA menampilkan dan mengekspor data mentah
dan TIDAK melakukan kalkulasi hari kerja, alfa, maupun potongan.
"""
from typing import Dict, Any, List, Optional, Tuple
from datetime import date, datetime
from sqlalchemy.orm import Session
from sqlalchemy import or_, desc, asc, func
from database.models import AttendanceRaw, Employee
from database.connection import get_db_session
from utils.logger import get_logger

logger = get_logger("AttendanceRawService")


class AttendanceRawService:
    """Layanan pengelolaan data absensi mentah."""

    @staticmethod
    def get_raw_records(
        db: Session,
        keyword: Optional[str] = None,
        month: Optional[int] = None,
        year: Optional[int] = None,
        start_date: Optional[date] = None,
        end_date: Optional[date] = None,
        batch_id: Optional[str] = None,
        employee_id: Optional[int] = None,
        sort_by: str = "tanggal",
        sort_dir: str = "asc",
        page: int = 1,
        page_size: int = 50,
    ) -> Tuple[List[Dict[str, Any]], int, int]:
        """
        Mengambil daftar transaksi absensi mentah dengan multi-filter dan paginasi.

        Returns:
            (records_list, total_records, total_pages)
        """
        query = db.query(AttendanceRaw)

        # 1. Filter Kata Kunci (Nama, Emp Num, No ID, NIK)
        if keyword and keyword.strip():
            kw = f"%{keyword.strip()}%"
            query = query.filter(
                or_(
                    AttendanceRaw.nama.ilike(kw),
                    AttendanceRaw.emp_num.ilike(kw),
                    AttendanceRaw.no_id.ilike(kw),
                    AttendanceRaw.nik.ilike(kw),
                )
            )

        # 2. Filter Batch ID
        if batch_id and batch_id.strip() and batch_id != "SEMUA":
            query = query.filter(AttendanceRaw.import_batch_id == batch_id.strip())

        # 3. Filter Employee ID
        if employee_id:
            query = query.filter(AttendanceRaw.employee_id == employee_id)

        # 4. Filter Rentang Tanggal Spesifik
        if start_date:
            query = query.filter(AttendanceRaw.tanggal >= start_date)
        if end_date:
            query = query.filter(AttendanceRaw.tanggal <= end_date)

        # 5. Filter Bulan dan Tahun (jika rentang tanggal tidak ditentukan)
        if not start_date and not end_date:
            if year:
                # Filter tahun
                query = query.filter(func.strftime("%Y", AttendanceRaw.tanggal) == f"{year:04d}")
            if month:
                # Filter bulan
                query = query.filter(func.strftime("%m", AttendanceRaw.tanggal) == f"{month:02d}")

        # Total record terhitung
        total_records = query.count()

        # 6. Pengurutan (Sorting)
        sort_column = getattr(AttendanceRaw, sort_by, AttendanceRaw.tanggal)
        if sort_dir.lower() == "desc":
            query = query.order_by(desc(sort_column), desc(AttendanceRaw.id))
        else:
            query = query.order_by(asc(sort_column), asc(AttendanceRaw.id))

        # 7. Paginasi
        total_pages = max(1, (total_records + page_size - 1) // page_size)
        offset = (page - 1) * page_size
        results = query.offset(offset).limit(page_size).all()

        records_data = []
        for idx, rec in enumerate(results, start=offset + 1):
            item = rec.to_dict()
            item["no"] = idx
            records_data.append(item)

        return records_data, total_records, total_pages

    @staticmethod
    def get_available_batches(db: Session) -> List[str]:
        """Mengambil seluruh daftar batch ID yang pernah tercatat di data raw."""
        batches = (
            db.query(AttendanceRaw.import_batch_id)
            .filter(AttendanceRaw.import_batch_id.isnot(None))
            .distinct()
            .order_by(desc(AttendanceRaw.import_batch_id))
            .all()
        )
        return [b[0] for b in batches if b[0]]

    @staticmethod
    def get_summary_statistics(db: Session, month: Optional[int] = None, year: Optional[int] = None) -> Dict[str, Any]:
        """Menghitung ringkasan statistik data mentah pada periode tertentu."""
        query = db.query(AttendanceRaw)
        if year:
            query = query.filter(func.strftime("%Y", AttendanceRaw.tanggal) == f"{year:04d}")
        if month:
            query = query.filter(func.strftime("%m", AttendanceRaw.tanggal) == f"{month:02d}")

        total_transaksi = query.count()
        total_karyawan_unik = query.with_entities(AttendanceRaw.nama).distinct().count()

        # Transaksi dengan scan masuk kosong
        jam_masuk_kosong = query.filter(or_(AttendanceRaw.scan_masuk.is_(None), AttendanceRaw.scan_masuk == "")).count()
        # Transaksi dengan scan pulang kosong
        jam_pulang_kosong = query.filter(or_(AttendanceRaw.scan_pulang.is_(None), AttendanceRaw.scan_pulang == "")).count()
        # Transaksi tanpa relasi employee
        unlinked = query.filter(AttendanceRaw.employee_id.is_(None)).count()

        return {
            "total_transaksi": total_transaksi,
            "total_karyawan_unik": total_karyawan_unik,
            "jam_masuk_kosong": jam_masuk_kosong,
            "jam_pulang_kosong": jam_pulang_kosong,
            "unlinked_count": unlinked,
        }

    @staticmethod
    def export_raw_to_excel(
        db: Session,
        filepath: str,
        keyword: Optional[str] = None,
        month: Optional[int] = None,
        year: Optional[int] = None,
        start_date: Optional[date] = None,
        end_date: Optional[date] = None,
        batch_id: Optional[str] = None
    ) -> str:
        """
        Mengekspor seluruh data transaksi absensi mentah hasil filter ke berkas Excel .xlsx.
        Menggunakan openpyxl dengan fallback CSV bertanda titik koma (;).
        """
        records, total, _ = AttendanceRawService.get_raw_records(
            db=db,
            keyword=keyword,
            month=month,
            year=year,
            start_date=start_date,
            end_date=end_date,
            batch_id=batch_id,
            page=1,
            page_size=100000,  # Ambil seluruh data hasil filter
        )

        headers = [
            "No", "Emp Num.", "No. ID.", "NIK", "Nama Karyawan",
            "Tanggal", "Scan Masuk", "Scan Pulang", "Source File",
            "Batch ID", "Waktu Import"
        ]

        try:
            import openpyxl
            from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Data Absensi Raw"

            # Header Styling
            header_font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
            header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
            thin_border = Border(
                left=Side(style="thin", color="CBD5E1"),
                right=Side(style="thin", color="CBD5E1"),
                top=Side(style="thin", color="CBD5E1"),
                bottom=Side(style="thin", color="CBD5E1")
            )

            ws.append(headers)
            for col_idx in range(1, len(headers) + 1):
                cell = ws.cell(row=1, column=col_idx)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = Alignment(horizontal="center", vertical="center")

            # Isi baris data
            for r in records:
                row_vals = [
                    r["no"],
                    r["emp_num"],
                    r["no_id"],
                    r["nik"],
                    r["nama"],
                    r["tanggal"],
                    r["scan_masuk"],
                    r["scan_pulang"],
                    r["source_file"],
                    r["import_batch_id"],
                    r["created_at"],
                ]
                ws.append(row_vals)
                row_idx = ws.max_row
                for c_idx in range(1, len(row_vals) + 1):
                    ws.cell(row=row_idx, column=c_idx).border = thin_border

            # Auto fit width
            for col in ws.columns:
                max_len = max(len(str(cell.value or "")) for cell in col)
                col_letter = openpyxl.utils.get_column_letter(col[0].column)
                ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

            wb.save(filepath)
            return filepath

        except ImportError:
            # Fallback CSV jika openpyxl belum terinstal
            import csv
            csv_path = filepath if filepath.endswith(".csv") else filepath.rsplit(".", 1)[0] + ".csv"
            with open(csv_path, "w", newline="", encoding="utf-8-sig") as f:
                writer = csv.writer(f, delimiter=";")
                writer.writerow(headers)
                for r in records:
                    writer.writerow([
                        r["no"],
                        r["emp_num"],
                        r["no_id"],
                        r["nik"],
                        r["nama"],
                        r["tanggal"],
                        r["scan_masuk"],
                        r["scan_pulang"],
                        r["source_file"],
                        r["import_batch_id"],
                        r["created_at"],
                    ])
            return csv_path
