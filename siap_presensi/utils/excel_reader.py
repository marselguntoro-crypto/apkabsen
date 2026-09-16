"""
Modul Excel Reader untuk Aplikasi Desktop SIAP.
Membaca file Microsoft Excel (.xlsx), mendeteksi daftar worksheet,
dan mengekstrak data baris secara efisien dan andal.
Menggunakan openpyxl jika tersedia, dengan fallback pembaca XML ZIP bawaan
sehingga kompatibel di segala lingkungan Python offline.
"""
import os
import zipfile
import xml.etree.ElementTree as ET
from pathlib import Path
from typing import List, Dict, Any, Tuple, Optional
from utils.logger import get_logger

logger = get_logger("ExcelReader")


def get_excel_sheet_names(filepath: str) -> List[str]:
    """
    Mengambil daftar nama worksheet yang ada di dalam berkas .xlsx.
    """
    path = Path(filepath)
    if not path.exists():
        raise FileNotFoundError(f"File '{filepath}' tidak ditemukan.")

    if not path.name.lower().endswith(".xlsx"):
        raise ValueError("Format file harus berupa Excel .xlsx.")

    # 1. Coba gunakan openpyxl jika terpasang
    try:
        import openpyxl
        wb = openpyxl.load_workbook(filepath, read_only=True, data_only=True)
        sheet_names = wb.sheetnames
        wb.close()
        return sheet_names
    except ImportError:
        logger.debug("openpyxl tidak terpasang, beralih ke fallback zipfile XML parser.")
    except Exception as e:
        logger.warning(f"openpyxl gagal membuka file: {e}. Menggunakan fallback parser.")

    # 2. Fallback parser menggunakan modul standar zipfile dan XML
    try:
        with zipfile.ZipFile(filepath, "r") as z:
            if "xl/workbook.xml" not in z.namelist():
                raise ValueError("Berkas .xlsx tidak valid atau rusak (xl/workbook.xml tidak ditemukan).")

            tree = ET.fromstring(z.read("xl/workbook.xml"))
            sheets = []
            for elem in tree.iter():
                if elem.tag.endswith("sheet") and "name" in elem.attrib:
                    sheets.append(elem.attrib["name"])

            if not sheets:
                sheets = ["Sheet1"]
            return sheets
    except Exception as err:
        raise ValueError(f"Gagal membaca worksheet Excel: {err}")


def read_excel_rows(filepath: str, sheet_name: Optional[str] = None) -> Tuple[List[str], List[Dict[str, Any]], int]:
    """
    Membaca seluruh baris dari sheet tertentu pada file .xlsx.

    Returns:
        (headers, rows_as_dicts, total_row_count)
        - headers: list string nama kolom header baris pertama
        - rows_as_dicts: list dictionary baris data {header_name: cell_value}
        - total_row_count: total baris data (di luar header)
    """
    path = Path(filepath)
    if not path.exists():
        raise FileNotFoundError(f"File '{filepath}' tidak ditemukan.")

    # 1. Gunakan openpyxl jika tersedia
    try:
        import openpyxl
        wb = openpyxl.load_workbook(filepath, read_only=False, data_only=True)
        if sheet_name and sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
        else:
            ws = wb.active

        raw_rows = list(ws.iter_rows(values_only=True))
        wb.close()

        if not raw_rows:
            return [], [], 0

        # Cari baris header (baris pertama yang tidak kosong)
        header_row_idx = 0
        while header_row_idx < len(raw_rows) and all(c is None or str(c).strip() == "" for c in raw_rows[header_row_idx]):
            header_row_idx += 1

        if header_row_idx >= len(raw_rows):
            return [], [], 0

        raw_headers = raw_rows[header_row_idx]
        headers = [str(h).strip() if h is not None else f"Column_{i+1}" for i, h in enumerate(raw_headers)]

        data_rows = []
        for row in raw_rows[header_row_idx + 1:]:
            # Lewati baris yang seluruh kolomnya kosong
            if all(c is None or str(c).strip() == "" for c in row):
                continue
            row_dict = {}
            for i, val in enumerate(row):
                if i < len(headers):
                    row_dict[headers[i]] = val
            data_rows.append(row_dict)

        return headers, data_rows, len(data_rows)

    except ImportError:
        logger.debug("openpyxl tidak terpasang, membaca xlsx via pure python fallback.")
    except Exception as e:
        logger.warning(f"openpyxl gagal memproses rows: {e}. Menggunakan pure python fallback.")

    # 2. Pure Python Fallback untuk XLSX
    return _read_xlsx_pure_python(filepath, sheet_name)


def _read_xlsx_pure_python(filepath: str, sheet_name: Optional[str] = None) -> Tuple[List[str], List[Dict[str, Any]], int]:
    """Fallback parser membaca XML di dalam arsip ZIP .xlsx."""
    with zipfile.ZipFile(filepath, "r") as z:
        # Baca shared strings
        shared_strings = []
        if "xl/sharedStrings.xml" in z.namelist():
            ss_root = ET.fromstring(z.read("xl/sharedStrings.xml"))
            for si in ss_root.findall(".//{http://schemas.openxmlformats.org/spreadsheetml/2006/main}si"):
                t = si.find("{http://schemas.openxmlformats.org/spreadsheetml/2006/main}t")
                if t is not None and t.text:
                    shared_strings.append(t.text)
                else:
                    # Gabungkan potongan teks jika formatted
                    texts = [elem.text for elem in si.findall(".//{http://schemas.openxmlformats.org/spreadsheetml/2006/main}t") if elem.text]
                    shared_strings.append("".join(texts))

        # Tentukan sheet xml
        sheet_target = "xl/worksheets/sheet1.xml"
        if "xl/workbook.xml" in z.namelist():
            wb_root = ET.fromstring(z.read("xl/workbook.xml"))
            ns = {"ns": "http://schemas.openxmlformats.org/spreadsheetml/2006/main", "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships"}
            sheets_elem = wb_root.find("ns:sheets", ns)
            if sheets_elem is not None:
                sheet_list = sheets_elem.findall("ns:sheet", ns)
                selected_idx = 1
                if sheet_name:
                    for idx, s in enumerate(sheet_list, 1):
                        if s.attrib.get("name") == sheet_name:
                            selected_idx = idx
                            break
                candidate = f"xl/worksheets/sheet{selected_idx}.xml"
                if candidate in z.namelist():
                    sheet_target = candidate

        if sheet_target not in z.namelist():
            # Temukan sheet pertama apa saja
            ws_files = [f for f in z.namelist() if f.startswith("xl/worksheets/sheet") and f.endswith(".xml")]
            if ws_files:
                sheet_target = ws_files[0]
            else:
                return [], [], 0

        sheet_root = ET.fromstring(z.read(sheet_target))
        ns = {"ns": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
        sheet_data = sheet_root.find("ns:sheetData", ns)
        if sheet_data is None:
            return [], [], 0

        rows = []
        for row in sheet_data.findall("ns:row", ns):
            row_vals = []
            for c in row.findall("ns:c", ns):
                val_type = c.attrib.get("t")
                v_elem = c.find("ns:v", ns)
                val = None
                if v_elem is not None and v_elem.text is not None:
                    raw_text = v_elem.text
                    if val_type == "s":
                        try:
                            s_idx = int(raw_text)
                            val = shared_strings[s_idx] if s_idx < len(shared_strings) else raw_text
                        except Exception:
                            val = raw_text
                    elif val_type == "b":
                        val = (raw_text == "1")
                    else:
                        val = raw_text
                row_vals.append(val)
            if any(v is not None and str(v).strip() != "" for v in row_vals):
                rows.append(row_vals)

        if not rows:
            return [], [], 0

        raw_headers = rows[0]
        headers = [str(h).strip() if h is not None else f"Column_{i+1}" for i, h in enumerate(raw_headers)]

        data_rows = []
        for row in rows[1:]:
            row_dict = {}
            for i, val in enumerate(row):
                if i < len(headers):
                    row_dict[headers[i]] = val
            data_rows.append(row_dict)

        return headers, data_rows, len(data_rows)
